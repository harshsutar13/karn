"""
tools/pptx_writer.py

Generate .pptx presentations from structured slide specs.
"""

from __future__ import annotations

import os
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt


def write_pptx(
    title: str,
    output_path: str,
    slides: list[dict],
    subtitle: str | None = None,
) -> str:
    """
    Write a PowerPoint presentation.

    Parameters
    ----------
    title : str
        Deck title (used on the first title slide).
    output_path : str
        Where to save the .pptx.
    slides : list[dict]
        Each dict has:
          - title (str)
          - bullets: list[str] (optional)
          - notes: str (optional) — speaker notes
    subtitle : str | None
        Subtitle for the title slide.

    Returns
    -------
    str : absolute path written
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- Title slide ---
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title
    subtitle_placeholder = slide.placeholders[1]
    subtitle_text = subtitle or f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    subtitle_placeholder.text = subtitle_text

    # --- Content slides ---
    content_layout = prs.slide_layouts[1]  # Title and Content

    for spec in slides:
        stitle = spec.get("title", "")
        bullets = spec.get("bullets") or []
        notes = spec.get("notes")

        slide = prs.slides.add_slide(content_layout)

        # Title
        if slide.shapes.title is not None:
            slide.shapes.title.text = stitle

        # Body
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        if bullets:
            for i, b in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = str(b)
                p.font.size = Pt(18)
        else:
            # Blank content slide
            tf.text = "(no bullets)"

        # Speaker notes
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.save(output_path)
    return os.path.abspath(output_path)