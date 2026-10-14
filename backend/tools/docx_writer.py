"""
tools/docx_writer.py

Generate a .docx deliverable from structured sections.
"""

from __future__ import annotations

import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor


def write_docx(
    title: str,
    sections: list[dict],
    output_path: str,
    metadata: dict | None = None,
    sources: list[str] | None = None,
    footer_note: str | None = None,
) -> str:
    """
    Write a Word document.

    Parameters
    ----------
    title : str
        Document title (Heading 0).
    sections : list[dict]
        Each dict: {"heading": str, "body": str}
    output_path : str
        Absolute path to save the .docx.
    metadata : dict | None
        Key-value pairs rendered as a table under the title.
    sources : list[str] | None
        Filenames cited, added as a "Sources" section.
    footer_note : str | None
        Small italic line at the bottom.

    Returns
    -------
    str : absolute path written
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = Document()

    # Title
    doc.add_heading(title, level=0)

    # Generated-by line
    p = doc.add_paragraph()
    run = p.add_run(
        f"AI-assisted draft · Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} · "
        f"PENDING HUMAN REVIEW"
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Metadata table
    if metadata:
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for key, val in metadata.items():
            row = table.add_row().cells
            row[0].text = str(key)
            row[1].text = str(val)
        doc.add_paragraph()

    # Sections
    for s in sections:
        heading = s.get("heading") or ""
        body = s.get("body") or ""
        if heading:
            doc.add_heading(heading, level=1)
        for para in body.split("\n\n"):
            para = para.strip()
            if para:
                doc.add_paragraph(para)

    # Sources
    if sources:
        doc.add_heading("Sources", level=1)
        for src in sources:
            doc.add_paragraph(src, style="List Bullet")

    # Footer note
    if footer_note:
        p = doc.add_paragraph()
        run = p.add_run(footer_note)
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.save(output_path)
    return os.path.abspath(output_path)